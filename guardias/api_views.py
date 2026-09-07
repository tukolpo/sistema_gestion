from datetime import datetime, date

from rest_framework.authentication import SessionAuthentication
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework_simplejwt.authentication import JWTAuthentication

from usuarios.constants import NIVEL_GESTION_USUARIOS
from usuarios.permissions import TieneJerarquiaMinima

from guardias.serializers import DisponibilidadFuncionarioSerializer
from guardias.services import consultar_disponibilidad


class PermisoGestionGuardias(TieneJerarquiaMinima):
    nivel_requerido = NIVEL_GESTION_USUARIOS


class DisponibilidadAPIView(APIView):
    """Consulta la disponibilidad de funcionarios desde la base de datos."""

    permission_classes = [IsAuthenticated, PermisoGestionGuardias]
    authentication_classes = [SessionAuthentication, JWTAuthentication]

    def get(self, request):
        fecha_param = request.query_params.get("fecha")
        fecha = None
        if fecha_param:
            try:
                fecha = datetime.strptime(fecha_param, "%Y-%m-%d").date()
            except ValueError:
                return Response(
                    {"detail": "Formato de fecha inválido. Use YYYY-MM-DD."},
                    status=400,
                )

        datos = consultar_disponibilidad(fecha)
        serializer = DisponibilidadFuncionarioSerializer(datos, many=True)
        return Response(serializer.data)

from rest_framework import status
from django.core.exceptions import ValidationError
from datetime import timedelta, datetime, date
import csv
from django.http import HttpResponse
from guardias.models import GuardiaTurno
from guardias.services import asignar_turno_funcionario

class AsignarTurnoAPIView(APIView):
    permission_classes = [IsAuthenticated, PermisoGestionGuardias]
    authentication_classes = [SessionAuthentication, JWTAuthentication]

    def post(self, request):
        trabajador_id = request.data.get("funcionario")
        fecha_str = request.data.get("fecha")
        turno = request.data.get("turno")

        try:
            fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date()
            guardia = asignar_turno_funcionario(trabajador_id, fecha, turno)
            return Response({"detail": "Turno asignado exitosamente.", "id": guardia.id}, status=201)
        except ValidationError as e:
            msg = e.messages[0] if hasattr(e, 'messages') else str(e)
            return Response({"detail": msg}, status=400)
        except Exception as e:
            return Response({"detail": str(e)}, status=400)


class CronogramaSemanalAPIView(APIView):
    permission_classes = [IsAuthenticated, PermisoGestionGuardias]
    authentication_classes = [SessionAuthentication, JWTAuthentication]

    def get(self, request):
        fecha_str = request.query_params.get("fecha")
        ref_fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date() if fecha_str else date.today()
        
        inicio_semana = ref_fecha - timedelta(days=ref_fecha.weekday())
        fin_semana = inicio_semana + timedelta(days=6)

        guardias = GuardiaTurno.objects.filter(fecha__range=[inicio_semana, fin_semana]).select_related("trabajador")
        
        data = [{
            "id": g.id,
            "trabajador_id": g.trabajador_id,
            "trabajador": str(g.trabajador),
            "fecha": str(g.fecha),
            "turno": g.turno,
            "aprobado": g.aprobado
        } for g in guardias]

        return Response({"inicio": str(inicio_semana), "fin": str(fin_semana), "guardias": data})

    def patch(self, request):
        fecha_str = request.data.get("fecha")
        aprobado = request.data.get("aprobado", True)
        
        ref_fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date() if fecha_str else date.today()
        inicio_semana = ref_fecha - timedelta(days=ref_fecha.weekday())
        fin_semana = inicio_semana + timedelta(days=6)

        GuardiaTurno.objects.filter(fecha__range=[inicio_semana, fin_semana]).update(aprobado=aprobado)
        return Response({"detail": "Cronograma aprobado correctamente."})


import openpyxl
from openpyxl.utils import get_column_letter

class ExportarReporteGuardiasView(APIView):
    permission_classes = [IsAuthenticated, PermisoGestionGuardias]
    authentication_classes = [SessionAuthentication, JWTAuthentication]

    def get(self, request):
        fecha_str = request.query_params.get("fecha")
        ref_fecha = datetime.strptime(fecha_str, "%Y-%m-%d").date() if fecha_str else date.today()
        inicio_semana = ref_fecha - timedelta(days=ref_fecha.weekday())
        fin_semana = inicio_semana + timedelta(days=6)

        guardias = GuardiaTurno.objects.filter(fecha__range=[inicio_semana, fin_semana]).select_related("trabajador")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Guardias"
        ws.append(["Fecha", "Funcionario", "Cédula", "Turno", "Estado"])

        for g in guardias:
            ws.append([
                str(g.fecha),
                str(g.trabajador),
                g.trabajador.cedula,
                g.get_turno_display(),
                "Aprobado" if g.aprobado else "Pendiente",
            ])

        for i, columna in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = 20

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="reporte_guardias_{inicio_semana}_al_{fin_semana}.xlsx"'
        wb.save(response)
        return response